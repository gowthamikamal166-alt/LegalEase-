"""Formatting utilities: sanitize, DOCX, PDF and HTML preview generators."""
import html
import io
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from fpdf import FPDF

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from config import DOC_LOGO_PATH, FOOTER_TEXT  # noqa: E402

_REPLACEMENTS = {
    "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
    "\u2013": "-", "\u2014": "-", "\u2026": "...", "\u2022": "-", "\u00a0": " ",
}


def sanitize_text(text: str) -> str:
    """Replace typographic characters and strip markdown emphasis markers."""
    for old, new in _REPLACEMENTS.items():
        text = text.replace(old, new)
    text = text.replace("**", "").replace("__", "")
    return text.strip()


def _is_heading(line: str) -> bool:
    s = line.strip()
    if s.startswith("#"):
        return True
    if len(s) >= 80:
        return False
    if re.match(r"^\d+\.\s.*:$", s):          # "1. Services:"
        return True
    return s.isupper() and len(s) > 3          # "WITNESSETH:" / "BETWEEN"


def _clean_heading(line: str) -> str:
    return line.strip().lstrip("#").strip()


def _split_terms(terms: str):
    return [t.strip() for t in (terms or "").split(";") if t.strip()]


def _latin(s: str) -> str:
    return s.encode("latin-1", "replace").decode("latin-1")


# --------------------------------------------------------------------- DOCX
def format_docx(text: str, doc_type: str, terms: str = "") -> bytes:
    text = sanitize_text(text)
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    if os.path.exists(DOC_LOGO_PATH):
        doc.add_picture(DOC_LOGO_PATH, width=Inches(2.2))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.title())
    run.bold = True
    run.font.size = Pt(16)

    for line in text.splitlines():
        s = line.strip()
        if not s:
            continue
        if _is_heading(s):
            doc.add_paragraph().add_run(_clean_heading(s)).bold = True
        elif s.startswith(("- ", "* ")):
            doc.add_paragraph(s[2:], style="List Bullet")
        else:
            doc.add_paragraph(s).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    items = _split_terms(terms)
    if items:
        doc.add_paragraph().add_run("Summary of Key Terms").bold = True
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text = "No.", "Term / Condition"
        for c in hdr:
            for r in c.paragraphs[0].runs:
                r.bold = True
        for i, t in enumerate(items, 1):
            row = table.add_row().cells
            row[0].text, row[1].text = str(i), t

    footer_p = doc.sections[0].footer.paragraphs[0]
    footer_p.text = FOOTER_TEXT
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in footer_p.runs:
        r.italic = True
        r.font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------- PDF
class _LegalPDF(FPDF):
    def __init__(self, doc_type: str):
        super().__init__()
        self.doc_type = doc_type

    def header(self):
        if os.path.exists(DOC_LOGO_PATH):
            w = 40
            self.image(DOC_LOGO_PATH, x=(self.w - w) / 2, y=8, w=w)
            self.set_y(26)
        else:
            self.set_y(10)
        self.set_font("Helvetica", "B", 13)
        self.cell(0, 8, _latin(self.doc_type.title()), align="C", new_x="LMARGIN", new_y="NEXT")
        self.ln(4)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 8, _latin(FOOTER_TEXT), align="C")


def format_pdf(text: str, doc_type: str) -> bytes:
    text = sanitize_text(text)
    pdf = _LegalPDF(doc_type)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    for line in text.splitlines():
        s = line.strip()
        if not s:
            pdf.ln(2)
            continue
        if _is_heading(s):
            pdf.set_font("Helvetica", "B", 11)
            pdf.multi_cell(0, 7, _latin(_clean_heading(s)), new_x="LMARGIN", new_y="NEXT")
        elif s.startswith(("- ", "* ")):
            pdf.set_font("Helvetica", "", 10.5)
            pdf.set_x(pdf.l_margin + 6)
            pdf.multi_cell(0, 6, _latin("- " + s[2:]), new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.set_font("Helvetica", "", 10.5)
            pdf.multi_cell(0, 6, _latin(s), align="J", new_x="LMARGIN", new_y="NEXT")

    out = pdf.output()
    return out.encode("latin-1") if isinstance(out, str) else bytes(out)


# --------------------------------------------------------------------- HTML
def format_html_preview(text: str) -> str:
    text = sanitize_text(text)
    parts = []
    for line in text.splitlines():
        s = line.strip()
        if not s:
            continue
        if _is_heading(s):
            parts.append(f"<h4 style='margin:18px 0 6px;color:#fff;'>{html.escape(_clean_heading(s))}</h4>")
        elif s.startswith(("- ", "* ")):
            parts.append(f"<li style='margin-left:20px;'>{html.escape(s[2:])}</li>")
        else:
            parts.append(f"<p style='margin:6px 0;line-height:1.6;'>{html.escape(s)}</p>")
    return "\n".join(parts)
